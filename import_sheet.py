# Imports the persona sheet (xlsx) into data.json: D.P = live library (exactly the sheet), D.R = retired personas kept only so old discussions still open.
import openpyxl, json, sys, shutil, time
from pathlib import Path
SOURCE = Path(__file__).resolve().parent
X = sys.argv[1]; DJ = SOURCE / 'data.json'
shutil.copy(DJ, SOURCE / f'data.backup-{time.strftime("%Y%m%d-%H%M%S")}.json')
D = json.load(open(DJ))
ws = openpyxl.load_workbook(X, data_only=True)['Personas']
H = [c.value for c in ws[1]]; col = {h: i for i, h in enumerate(H)}
def g(r, h): v = r[col[h]]; return "" if v is None or str(v).strip() in ("", "None") and h not in ("Sensitivity_Level",) else str(v).strip()
P = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    doms = [x.strip() for x in g(r, 'Domain').split(';') if x.strip()]
    subs = [x.strip() for x in g(r, 'Sub_Domain').split(';') if x.strip()]
    P.append([g(r,'Persona_ID'), g(r,'Persona_Name'), g(r,'Persona_Name_Arabic'), g(r,'Basis_Type'), doms[0], subs[0] if subs else "",
              g(r,'Era'), g(r,'Region_Culture'), g(r,'Tagline'), g(r,'Brief_Description'), g(r,'Worldview_Core_Beliefs'), g(r,'Reasoning_Style'),
              g(r,'Blind_Spots'), g(r,'Key_Values'), g(r,'Communication_Tone'), g(r,'Sensitivity_Level') or "None",
              g(r,'Typical_Arguments'), g(r,'Signature_Questions'), g(r,'Natural_Opponent_IDs'), g(r,'Best_Suited_Topics'),
              g(r,'Sensitivity_Guidance'), g(r,'Gender'), '; '.join(doms), '; '.join(subs), g(r,'Source_References')])
ids = {p[0] for p in P}
assert len(ids) == len(P), 'duplicate IDs'
old = list(D['P']) + list(D.get('R', []))
R = [(o[:16] + [""] * 9)[:25] for o in old if o[0] not in ids]
R = list({o[0]: o for o in R}.values())
D['P'] = P; D['R'] = R
json.dump(D, open(DJ, 'w'), ensure_ascii=False)
print('library', len(P), 'retired (old discussions only)', len(R))
